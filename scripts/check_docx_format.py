#!/usr/bin/env python3
"""Generate a lightweight submission format report for a Word document."""

from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path
from zipfile import ZipFile

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError as exc:  # pragma: no cover - dependency check path
    raise SystemExit(
        "python-docx is required. See references/install-dependencies.md"
    ) from exc

from sanitize_reference_docx import inspect_docx_structure

HEADING_STYLE_NAMES = {"Title", "Heading 1", "Heading 2", "Heading 3", "标题 1", "标题 2", "标题 3"}
TITLE_STYLE_NAMES = {"Title", "题名"}
LEVEL_ONE_HEADING_STYLE_NAMES = {"Heading 1", "标题 1"}
SECONDARY_HEADING_STYLE_NAMES = {"Heading 2", "标题 2"}
TERTIARY_HEADING_STYLE_NAMES = {"Heading 3", "标题 3"}
LEADING_SPACE_RE = re.compile(r"^[\s\u3000\xa0]+")
CHINESE_TERTIARY_HEADING_RE = re.compile(r"^[（(]\s*[一二三四五六七八九十]+\s*[）)]")
ARABIC_BODY_HEADING_RE = re.compile(r"^(?:[（(]\s*)?\d{1,2}\s*(?:[.．、）)]|\s+[^\d])")
LABEL_RE = re.compile(
    r"^(摘要|内容提要|关键词|作者|作者单位|作者简介|联系方式|基金项目|资助项目|通讯作者)\s*[:：]"
    r"|^(Abstract|Keywords?|Author|Corresponding author|Contact|Funding)\b\s*[:：.]?",
    re.IGNORECASE,
)
CAPTION_RE = re.compile(r"^(图|表)\s*[0-9０-９一二三四五六七八九十]+")


@dataclass
class Check:
    status: str
    item: str
    source: str
    detail: str


def paragraph_texts(doc: Document) -> list[str]:
    return [p.text.strip() for p in doc.paragraphs if p.text.strip()]


def full_text(texts: list[str]) -> str:
    return "\n".join(texts)


def has_pattern(text: str, patterns: list[str]) -> bool:
    return any(re.search(pattern, text, re.IGNORECASE) for pattern in patterns)


def first_body_paragraph(texts: list[str]) -> str:
    for text in texts:
        if len(text) >= 12 and not text.startswith("#"):
            return text
    return ""


def style_name(paragraph) -> str:
    return paragraph.style.name if paragraph.style else ""


def effective_first_line_indent(paragraph):
    indent = paragraph.paragraph_format.first_line_indent
    if indent is None and paragraph.style is not None:
        indent = paragraph.style.paragraph_format.first_line_indent
    return indent


def effective_left_indent(paragraph):
    indent = paragraph.paragraph_format.left_indent
    if indent is None and paragraph.style is not None:
        indent = paragraph.style.paragraph_format.left_indent
    return indent


def xml_indent_value(ppr, attr: str) -> str | None:
    if ppr is None:
        return None
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        return None
    return ind.get(qn(attr))


def effective_page_break_before(paragraph) -> bool:
    value = paragraph.paragraph_format.page_break_before
    if value is None and paragraph.style is not None:
        value = paragraph.style.paragraph_format.page_break_before
    return bool(value)


def effective_alignment(paragraph):
    alignment = paragraph.paragraph_format.alignment
    if alignment is None and paragraph.style is not None:
        alignment = paragraph.style.paragraph_format.alignment
    return alignment


def effective_keep_with_next(paragraph) -> bool:
    value = paragraph.paragraph_format.keep_with_next
    if value is None and paragraph.style is not None:
        value = paragraph.style.paragraph_format.keep_with_next
    return bool(value)


def effective_keep_together(paragraph) -> bool:
    value = paragraph.paragraph_format.keep_together
    if value is None and paragraph.style is not None:
        value = paragraph.style.paragraph_format.keep_together
    return bool(value)


def is_heading_paragraph(paragraph) -> bool:
    name = style_name(paragraph)
    return name in HEADING_STYLE_NAMES or name.startswith("Heading")


def is_label_or_caption(text: str) -> bool:
    return bool(LABEL_RE.search(text.strip()) or CAPTION_RE.search(text.strip()))


def is_chinese_tertiary_heading(text: str) -> bool:
    return bool(CHINESE_TERTIARY_HEADING_RE.search(text.strip()))


def is_arabic_body_heading(text: str) -> bool:
    return bool(ARABIC_BODY_HEADING_RE.search(text.strip()))


def check_presence(text: str, item: str, patterns: list[str], source: str, detail: str) -> Check:
    if has_pattern(text, patterns):
        return Check("PASS", item, source, detail)
    return Check("FAIL", item, source, f"未检测到：{detail}")


def inspect_sections(doc: Document) -> list[Check]:
    checks: list[Check] = []
    if not doc.sections:
        return [Check("WARN", "页面设置", "内置模板建议", "无法读取 section 信息")]

    section = doc.sections[0]
    margins = [
        section.top_margin.cm,
        section.bottom_margin.cm,
        section.left_margin.cm,
        section.right_margin.cm,
    ]
    if all(1.5 <= margin <= 3.5 for margin in margins):
        checks.append(Check("PASS", "页边距", "内置模板建议", "页边距处于常见投稿范围"))
    else:
        checks.append(
            Check(
                "WARN",
                "页边距",
                "内置模板建议",
                "页边距偏离常见投稿范围，请按目标期刊模板复核",
            )
        )
    return checks


def inspect_styles(doc: Document) -> list[Check]:
    checks: list[Check] = []
    headings = [p for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
    if headings:
        checks.append(Check("PASS", "标题层级", "内置模板建议", f"检测到 {len(headings)} 个标题样式段落"))
    else:
        checks.append(Check("WARN", "标题层级", "内置模板建议", "未检测到 Word 标题样式，可能影响目录和审稿阅读"))

    normal = next((style for style in doc.styles if style.name == "Normal"), None)
    if normal and normal.font and normal.font.size:
        checks.append(Check("PASS", "正文样式", "内置模板建议", f"Normal 字号：{normal.font.size.pt:.1f} pt"))
    else:
        checks.append(Check("WARN", "正文样式", "内置模板建议", "未能读取 Normal 字号，请人工复核"))

    page_break_styles = []
    keep_styles = []
    for style_name in ["Title", "Heading 1", "Heading 2", "Heading 3", "标题 1", "标题 2", "标题 3"]:
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        if style.paragraph_format.page_break_before:
            page_break_styles.append(style_name)
        if style.paragraph_format.keep_with_next or style.paragraph_format.keep_together:
            keep_styles.append(style_name)
    if page_break_styles:
        checks.append(
            Check(
                "FAIL",
                "标题分页设置",
                "导出清洗规则",
                "导出投稿稿默认不得保留标题自动分页：" + ", ".join(page_break_styles),
            )
        )
    else:
        checks.append(Check("PASS", "标题分页设置", "内置模板建议", "标题样式未设置自动分页"))
    if keep_styles:
        checks.append(
            Check(
                "WARN",
                "标题同页控制",
                "导出清洗规则",
                "标题样式仍带 keepNext/keepLines，可能造成异常空白或换页：" + ", ".join(keep_styles),
            )
        )
    else:
        checks.append(Check("PASS", "标题同页控制", "导出清洗规则", "标题样式未设置 keepNext/keepLines"))

    bad_level_one_styles = []
    for style_name in sorted(LEVEL_ONE_HEADING_STYLE_NAMES):
        try:
            style = doc.styles[style_name]
        except KeyError:
            continue
        pf = style.paragraph_format
        first = pf.first_line_indent.cm if pf.first_line_indent is not None else 0
        left = pf.left_indent.cm if pf.left_indent is not None else 0
        ppr = style._element.pPr
        first_chars = xml_indent_value(ppr, "w:firstLineChars")
        hanging_chars = xml_indent_value(ppr, "w:hangingChars")
        if (
            abs(first) > 0.01
            or abs(left) > 0.01
            or pf.alignment != WD_ALIGN_PARAGRAPH.CENTER
            or first_chars not in {None, "0"}
            or hanging_chars is not None
        ):
            bad_level_one_styles.append(
                f"{style_name}(first={first:.2f}cm,left={left:.2f}cm,"
                f"firstLineChars={first_chars},hangingChars={hanging_chars},align={pf.alignment})"
            )
    if bad_level_one_styles:
        checks.append(
            Check(
                "FAIL",
                "一级标题样式",
                "导出清洗规则",
                "Heading 1/标题 1 必须居中且无首行/左缩进：" + ", ".join(bad_level_one_styles),
            )
        )
    else:
        checks.append(Check("PASS", "一级标题样式", "导出清洗规则", "Heading 1/标题 1 居中且无首行/左缩进"))
    return checks


def inspect_page_break_xml(docx_path: Path) -> list[Check]:
    with ZipFile(docx_path) as archive:
        styles_xml = archive.read("word/styles.xml").decode("utf-8", errors="ignore")
        document_xml = archive.read("word/document.xml").decode("utf-8", errors="ignore")
    style_breaks = styles_xml.count("w:pageBreakBefore")
    explicit_breaks = document_xml.count('w:type="page"') + document_xml.count("w:pageBreakBefore")
    style_keep_controls = styles_xml.count("w:keepNext") + styles_xml.count("w:keepLines")
    document_keep_controls = document_xml.count("w:keepNext") + document_xml.count("w:keepLines")
    checks = []
    if style_breaks:
        checks.append(
            Check(
                "FAIL",
                "样式分页 XML",
                "导出清洗规则",
                f"styles.xml 中检测到 {style_breaks} 个 pageBreakBefore；导出稿默认应移除标题自动分页",
            )
        )
    else:
        checks.append(Check("PASS", "样式分页 XML", "内置模板建议", "styles.xml 无 pageBreakBefore"))
    if explicit_breaks:
        checks.append(Check("WARN", "文档分页 XML", "用户正文", f"document.xml 中检测到 {explicit_breaks} 个显式分页标记"))
    else:
        checks.append(Check("PASS", "文档分页 XML", "用户正文", "document.xml 无显式分页标记"))
    if style_keep_controls:
        checks.append(
            Check(
                "WARN",
                "样式同页 XML",
                "导出清洗规则",
                f"styles.xml 中检测到 {style_keep_controls} 个 keepNext/keepLines；若标题或题注仍异常换页，请清理模板样式",
            )
        )
    else:
        checks.append(Check("PASS", "样式同页 XML", "导出清洗规则", "styles.xml 无 keepNext/keepLines"))
    if document_keep_controls:
        checks.append(
            Check(
                "WARN",
                "段落同页 XML",
                "导出清洗规则",
                f"document.xml 中检测到 {document_keep_controls} 个 keepNext/keepLines，可能造成异常换页",
            )
        )
    else:
        checks.append(Check("PASS", "段落同页 XML", "导出清洗规则", "document.xml 无 keepNext/keepLines"))
    return checks


def inspect_content_artifacts(docx_path: Path) -> list[Check]:
    summary = inspect_docx_structure(docx_path)
    checks: list[Check] = []
    if summary["tables"]:
        checks.append(
            Check(
                "WARN",
                "表格对象",
                "内置模板建议",
                f"检测到 {summary['tables']} 个表格对象；请确认均为正文表格且按三线表排版",
            )
        )
    else:
        checks.append(Check("PASS", "表格对象", "内置模板建议", "未检测到模板残留表格对象"))

    visual_count = summary["media_files"] + summary["drawings"] + summary["inline_shapes"]
    if visual_count:
        checks.append(
            Check(
                "WARN",
                "图片/绘图对象",
                "内置模板建议",
                "检测到 "
                f"{summary['media_files']} 个媒体文件、{summary['drawings']} 个绘图对象、"
                f"{summary['inline_shapes']} 个行内图片；请确认不是模板残留",
            )
        )
    else:
        checks.append(Check("PASS", "图片/绘图对象", "内置模板建议", "未检测到图片或绘图残留"))

    if summary["max_empty_paragraph_run"] > 8:
        checks.append(
            Check(
                "WARN",
                "异常空段落",
                "内置模板建议",
                f"检测到连续 {summary['max_empty_paragraph_run']} 个空段落，请复核是否造成异常空白",
            )
        )
    else:
        checks.append(Check("PASS", "异常空段落", "内置模板建议", "未检测到长串连续空段落"))
    return checks


def inspect_markdown_export_artifacts(doc: Document) -> list[Check]:
    checks: list[Check] = []
    separator_paragraphs = [
        i
        for i, paragraph in enumerate(doc.paragraphs, start=1)
        if re.fullmatch(r"\s*[-*_]{3,}\s*", paragraph.text or "")
    ]
    if separator_paragraphs:
        checks.append(
            Check(
                "FAIL",
                "Markdown 分割线残留",
                "导出清洗规则",
                "以下段落仍为 Markdown 分割线：" + ", ".join(map(str, separator_paragraphs[:12])),
            )
        )
    else:
        checks.append(Check("PASS", "Markdown 分割线残留", "导出清洗规则", "未检测到独立 ---/***/___ 段落"))

    indented_headings = []
    for i, paragraph in enumerate(doc.paragraphs, start=1):
        style_name = paragraph.style.name if paragraph.style else ""
        if style_name not in {"Title", "Heading 1", "Heading 2", "Heading 3", "标题 1", "标题 2", "标题 3"} and not style_name.startswith("Heading"):
            continue
        indent = paragraph.paragraph_format.first_line_indent
        if indent is None and paragraph.style is not None:
            indent = paragraph.style.paragraph_format.first_line_indent
        if indent is not None and abs(indent.cm) > 0.01:
            indented_headings.append(f"{i}:{style_name}({indent.cm:.2f}cm)")
    if indented_headings:
        checks.append(
            Check(
                "FAIL",
                "标题首行缩进",
                "导出清洗规则",
                "标题/题名段落不得首行缩进：" + ", ".join(indented_headings[:12]),
            )
        )
    else:
        checks.append(Check("PASS", "标题首行缩进", "导出清洗规则", "题名和标题段落无首行缩进"))

    leading_space_headings = []
    for i, paragraph in enumerate(doc.paragraphs, start=1):
        if not is_heading_paragraph(paragraph):
            continue
        raw_text = paragraph.text or ""
        if LEADING_SPACE_RE.search(raw_text):
            leading_space_headings.append(f"{i}:{style_name(paragraph)}")
    if leading_space_headings:
        checks.append(
            Check(
                "FAIL",
                "标题文本前导空格",
                "导出清洗规则",
                "标题文本本身不得以空格或全角空格开头：" + ", ".join(leading_space_headings[:12]),
            )
        )
    else:
        checks.append(Check("PASS", "标题文本前导空格", "导出清洗规则", "标题文本无真实前导空格"))

    return checks


def inspect_heading_cleanup(doc: Document) -> list[Check]:
    checks: list[Check] = []
    first_nonempty = next((p for p in doc.paragraphs if p.text.strip()), None)
    if first_nonempty is None:
        checks.append(Check("FAIL", "主标题样式", "导出清洗规则", "未检测到非空段落"))
    else:
        first_style = style_name(first_nonempty)
        if first_style in LEVEL_ONE_HEADING_STYLE_NAMES:
            checks.append(Check("PASS", "主标题样式", "导出清洗规则", f"首个非空段落样式为 {first_style}"))
        elif first_style in TITLE_STYLE_NAMES:
            checks.append(Check("FAIL", "主标题样式", "导出清洗规则", "主标题不得使用 Title/题名样式，应使用 Heading 1/标题 1"))
        else:
            checks.append(
                Check(
                    "WARN",
                    "主标题样式",
                    "导出清洗规则",
                    f"首个非空段落样式为 {first_style or '未知'}，请人工确认是否为主标题",
                )
            )

    heading_breaks = []
    heading_keeps = []
    bad_level_one_paragraphs = []
    bad_secondary_heading_paragraphs = []
    bad_tertiary_marker_paragraphs = []
    bad_arabic_body_headings = []
    seen_first_heading = False
    for i, paragraph in enumerate(doc.paragraphs, start=1):
        name = style_name(paragraph)
        if is_heading_paragraph(paragraph):
            if is_arabic_body_heading(paragraph.text):
                bad_arabic_body_headings.append(f"{i}:{paragraph.text.strip()[:18]}({name})")
            if not seen_first_heading:
                seen_first_heading = True
            elif name not in SECONDARY_HEADING_STYLE_NAMES and name not in TERTIARY_HEADING_STYLE_NAMES:
                bad_secondary_heading_paragraphs.append(f"{i}:{paragraph.text.strip()[:18]}({name})")
            if is_chinese_tertiary_heading(paragraph.text) and name not in TERTIARY_HEADING_STYLE_NAMES:
                bad_tertiary_marker_paragraphs.append(f"{i}:{paragraph.text.strip()[:18]}({name})")
            if effective_page_break_before(paragraph):
                heading_breaks.append(f"{i}:{name}")
            if effective_keep_with_next(paragraph) or effective_keep_together(paragraph):
                heading_keeps.append(f"{i}:{name}")
            if name in LEVEL_ONE_HEADING_STYLE_NAMES:
                first_indent = effective_first_line_indent(paragraph)
                left_indent = effective_left_indent(paragraph)
                first_cm = first_indent.cm if first_indent is not None else 0
                left_cm = left_indent.cm if left_indent is not None else 0
                ppr = paragraph._p.pPr
                first_chars = xml_indent_value(ppr, "w:firstLineChars")
                hanging_chars = xml_indent_value(ppr, "w:hangingChars")
                if (
                    abs(first_cm) > 0.01
                    or abs(left_cm) > 0.01
                    or first_chars not in {None, "0"}
                    or hanging_chars is not None
                    or effective_alignment(paragraph) != WD_ALIGN_PARAGRAPH.CENTER
                ):
                    bad_level_one_paragraphs.append(
                        f"{i}:{paragraph.text.strip()[:18]}(first={first_cm:.2f}cm,left={left_cm:.2f}cm,"
                        f"firstLineChars={first_chars},hangingChars={hanging_chars})"
                    )
    if heading_breaks:
        checks.append(
            Check(
                "FAIL",
                "标题段落自动分页",
                "导出清洗规则",
                "标题段落不得设置 pageBreakBefore：" + ", ".join(heading_breaks[:12]),
            )
        )
    else:
        checks.append(Check("PASS", "标题段落自动分页", "导出清洗规则", "标题段落未设置自动分页"))
    if heading_keeps:
        checks.append(
            Check(
                "WARN",
                "标题段落同页控制",
                "导出清洗规则",
                "标题段落仍带 keepNext/keepLines，可能造成异常换页：" + ", ".join(heading_keeps[:12]),
            )
        )
    else:
        checks.append(Check("PASS", "标题段落同页控制", "导出清洗规则", "标题段落未设置 keepNext/keepLines"))
    if bad_level_one_paragraphs:
        checks.append(
            Check(
                "FAIL",
                "一级标题段落",
                "导出清洗规则",
                "Heading 1/标题 1 段落必须居中且无首行/左缩进：" + ", ".join(bad_level_one_paragraphs[:12]),
            )
        )
    else:
        checks.append(Check("PASS", "一级标题段落", "导出清洗规则", "Heading 1/标题 1 段落居中且无首行/左缩进"))
    if bad_secondary_heading_paragraphs:
        checks.append(
            Check(
                "FAIL",
                "其余标题层级",
                "导出清洗规则",
                "主标题之后的非三级标题必须使用 Heading 2/标题 2；原三级标题可保留 Heading 3/标题 3："
                + ", ".join(bad_secondary_heading_paragraphs[:12]),
            )
        )
    else:
        checks.append(
            Check(
                "PASS",
                "其余标题层级",
                "导出清洗规则",
                "主标题之后的标题为 Heading 2/标题 2，原三级标题保留 Heading 3/标题 3",
            )
        )
    if bad_tertiary_marker_paragraphs:
        checks.append(
            Check(
                "FAIL",
                "中文三级标题",
                "导出清洗规则",
                "`（一）...` 这类中文三级标题必须使用 Heading 3/标题 3："
                + ", ".join(bad_tertiary_marker_paragraphs[:12]),
            )
        )
    else:
        checks.append(Check("PASS", "中文三级标题", "导出清洗规则", "`（一）...` 层级标题均为 Heading 3/标题 3"))
    if bad_arabic_body_headings:
        checks.append(
            Check(
                "FAIL",
                "数字小标题正文化",
                "导出清洗规则",
                "`1. ...` 这类数字小标题应使用正文 Normal，不应保留标题样式："
                + ", ".join(bad_arabic_body_headings[:12]),
            )
        )
    else:
        checks.append(Check("PASS", "数字小标题正文化", "导出清洗规则", "`1. ...` 数字小标题未保留标题样式"))
    return checks


def inspect_label_indentation(doc: Document) -> list[Check]:
    bad_labels = []
    seen = 0
    for i, paragraph in enumerate(doc.paragraphs, start=1):
        text = paragraph.text.strip()
        raw_text = paragraph.text or ""
        if not text or not is_label_or_caption(raw_text):
            continue
        seen += 1
        first_indent = effective_first_line_indent(paragraph)
        left_indent = effective_left_indent(paragraph)
        first_cm = first_indent.cm if first_indent is not None else 0
        left_cm = left_indent.cm if left_indent is not None else 0
        has_leading_space = bool(LEADING_SPACE_RE.search(raw_text))
        if abs(first_cm) > 0.01 or abs(left_cm) > 0.01 or has_leading_space:
            preview = text[:18].replace("|", " ")
            text_space = ",text-leading-space" if has_leading_space else ""
            bad_labels.append(f"{i}:{preview}(first={first_cm:.2f}cm,left={left_cm:.2f}cm{text_space})")

    if bad_labels:
        return [
            Check(
                "FAIL",
                "标签段落缩进",
                "导出清洗规则",
                "摘要、关键词、作者信息和图表题不得继承正文缩进：" + ", ".join(bad_labels[:12]),
            )
        ]
    if seen:
        return [Check("PASS", "标签段落缩进", "导出清洗规则", f"检测到 {seen} 个标签/题注段落，均无首行或左缩进")]
    return [Check("PASS", "标签段落缩进", "导出清洗规则", "未检测到标签缩进风险")]


def inspect_sociological_research(text: str, texts: list[str]) -> list[Check]:
    checks = [
        Check(
            "PASS" if first_body_paragraph(texts) else "FAIL",
            "标题/正文",
            "官方要求",
            "检测稿件是否包含可读标题和正文段落",
        ),
        check_presence(text, "作者单位", [r"作者", r"单位", r"作者简介"], "官方要求", "作者姓名、单位或作者简介"),
        check_presence(text, "中文摘要", [r"摘要", r"内容提要"], "官方要求", "摘要或内容提要"),
        check_presence(text, "英文摘要", [r"\bAbstract\b"], "参照来源", "英文摘要 Abstract"),
        check_presence(text, "中文关键词", [r"关键词"], "官方要求", "中文关键词"),
        check_presence(text, "英文关键词", [r"\bKeywords?\b"], "参照来源", "英文关键词 Keywords"),
        check_presence(text, "注释", [r"注释", r"脚注", r"\[\^.+?\]"], "参照来源", "注释、脚注或 Markdown 脚注标记"),
        check_presence(text, "参考文献", [r"参考文献", r"References"], "官方要求", "参考文献标题"),
        check_presence(text, "作者简介", [r"作者简介"], "官方要求", "作者简介"),
        check_presence(text, "联系方式", [r"联系方式", r"邮箱", r"Email", r"E-mail"], "官方要求", "联系方式或邮箱"),
    ]
    return checks


def inspect_generic(text: str, style: str) -> list[Check]:
    source = "APA 体例" if style == "apa" else "GB/T 7714 体例"
    return [
        check_presence(text, "摘要", [r"摘要", r"Abstract"], source, "摘要"),
        check_presence(text, "关键词", [r"关键词", r"Keywords?"], source, "关键词"),
        check_presence(text, "参考文献", [r"参考文献", r"References"], source, "参考文献"),
    ]


def inspect_citations(text: str) -> list[Check]:
    citation_patterns = [
        r"[\u4e00-\u9fa5]{2,10}（\d{4}[a-z]?）",
        r"（[\u4e00-\u9fa5]{2,10}，\d{4}[a-z]?）",
        r"\b[A-Z][A-Za-z-]+ \(\d{4}[a-z]?\)",
        r"\([A-Z][A-Za-z-]+,\s*\d{4}[a-z]?\)",
        r"\[@[A-Za-z0-9_:-]+\]",
    ]
    if has_pattern(text, citation_patterns):
        return [Check("PASS", "正文引用", "引用规则", "检测到可识别的作者-年份或 Pandoc 引用")]
    return [Check("WARN", "正文引用", "引用规则", "未检测到常见正文引用形式")]


def render_report(checks: list[Check], docx: Path, style: str, reference_doc: str | None) -> str:
    lines = [
        "# Format Check Report",
        "",
        f"- Document: `{docx}`",
        f"- Style: `{style}`",
        "",
        "| Status | Item | Source | Detail |",
        "|---|---|---|---|",
    ]
    for check in checks:
        detail = check.detail.replace("|", "\\|")
        lines.append(f"| {check.status} | {check.item} | {check.source} | {detail} |")

    failed = sum(1 for check in checks if check.status == "FAIL")
    warned = sum(1 for check in checks if check.status == "WARN")
    lines.extend(
        [
            "",
            "## Summary",
            "",
            f"- FAIL: {failed}",
            f"- WARN: {warned}",
            "- 注：`官方要求`、`参照来源`、`内置模板建议` 的含义见 `references/journal-style-sociological-research.md`。",
        ]
    )
    return "\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description="Check a submission docx and write a Markdown report.")
    parser.add_argument("docx", help="Path to submission.docx")
    parser.add_argument(
        "--source-md",
        help="Optional source manuscript.md. Reserved for workflow traceability.",
    )
    parser.add_argument(
        "--style",
        choices=["sociological-research", "gb", "apa"],
        default="sociological-research",
    )
    parser.add_argument("--reference-doc", help="User-provided Word template used for export.")
    parser.add_argument(
        "--out",
        default="paper-workspace/06-submission/format-check-report.md",
        help="Output Markdown report path.",
    )
    args = parser.parse_args()

    docx_path = Path(args.docx).expanduser().resolve()
    if not docx_path.exists():
        raise SystemExit(f"docx not found: {docx_path}")

    doc = Document(str(docx_path))
    texts = paragraph_texts(doc)
    text = full_text(texts)
    checks: list[Check] = []

    if args.reference_doc:
        reference = Path(args.reference_doc).expanduser().resolve()
        if not reference.exists():
            raise SystemExit(f"reference doc not found: {reference}")
        checks.append(Check("PASS", "用户模板", "用户提供模板", "已使用用户提供的 Word 模板"))

    checks.extend(inspect_sections(doc))
    checks.extend(inspect_styles(doc))
    checks.extend(inspect_page_break_xml(docx_path))
    checks.extend(inspect_content_artifacts(docx_path))
    checks.extend(inspect_markdown_export_artifacts(doc))
    checks.extend(inspect_heading_cleanup(doc))
    checks.extend(inspect_label_indentation(doc))
    if args.style == "sociological-research":
        checks.extend(inspect_sociological_research(text, texts))
    else:
        checks.extend(inspect_generic(text, args.style))
    checks.extend(inspect_citations(text))

    out = Path(args.out).expanduser()
    if not out.is_absolute():
        out = (Path.cwd() / out).resolve()
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(render_report(checks, docx_path, args.style, args.reference_doc), encoding="utf-8")
    print(f"wrote: {out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
